"""Read-only CSV/XLSX quality report. No network, cleaning or metric inference."""
import argparse
import csv
import hashlib
import json
from collections import Counter
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path


def read_rows(source, encoding, sheet):
    warnings = []
    formulas = 0
    if source.suffix.lower() == '.csv':
        if sheet:
            raise ValueError('CSV 不接受 --sheet。')
        with source.open(encoding=encoding, newline='') as handle:
            rows = [row for row in csv.reader(handle, strict=True) if row]
        selected = None
    elif source.suffix.lower() == '.xlsx':
        try:
            import openpyxl
        except ImportError as error:
            raise ValueError('读取 Excel 需要技能 scripts/requirements.txt 中的 openpyxl。') from error
        raw = openpyxl.load_workbook(source, read_only=True, data_only=False)
        cached = openpyxl.load_workbook(source, read_only=True, data_only=True)
        try:
            selected = sheet or raw.sheetnames[0]
            if selected not in raw.sheetnames:
                raise ValueError('指定工作表不存在。')
            formula_rows = list(raw[selected].iter_rows(values_only=True))
            rows = list(cached[selected].iter_rows(values_only=True))
            formulas = sum(isinstance(v, str) and v.startswith('=') for row in formula_rows for v in row)
            if formulas:
                warnings.append('存在 Excel 公式：仅读取文件保存的缓存值，不能确认缓存是否最新；无缓存视为缺失。请在表格应用重算保存并复核。')
        finally:
            raw.close()
            cached.close()
    else:
        raise ValueError('仅支持 .csv 或 .xlsx；扫描文件、旧 .xls、宏文件需另处理。')
    if not rows:
        raise ValueError('文件为空。')
    headers = rows[0]
    if any(not isinstance(v, str) or not v.strip() for v in headers):
        raise ValueError('表头不能为空且必须为文本。')
    headers = [v.strip() for v in headers]
    if len(set(headers)) != len(headers):
        raise ValueError('表头重复（去除首尾空格后）。')
    data = rows[1:]
    for index, row in enumerate(data, 2):
        if len(row) != len(headers):
            raise ValueError(f'第 {index} 行列数不匹配。')
    return headers, data, selected, formulas, warnings


def missing(value):
    return value is None or (isinstance(value, str) and not value.strip())


def analyze(source, numeric, encoding='utf-8-sig', sheet=None):
    headers, rows, selected, formulas, warnings = read_rows(source, encoding, sheet)
    if len(set(numeric)) != len(numeric) or any(c not in headers for c in numeric):
        raise ValueError('数值列必须存在且不能重复指定。')
    # repr preserves types; no row content is included in the emitted report.
    duplicate_count = sum(count - 1 for count in Counter(tuple(repr(v) for v in row) for row in rows).values())
    report = {
        'schemaVersion': 1, 'sourceSha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'sheet': selected, 'rows': len(rows), 'columns': headers,
        'duplicateRowsBeyondFirst': duplicate_count, 'formulaCells': formulas,
        'missingByColumn': {h: sum(missing(row[i]) for row in rows) for i, h in enumerate(headers)},
        'numeric': {}, 'warnings': warnings,
        'limitations': ['只做质量和所选列统计，不清洗原表、不推定指标、不做因果分析。', '数值未解释单位；文本货币、分组逗号、百分号不自动转换。']
    }
    for column in numeric:
        index = headers.index(column)
        values, absent, invalid = [], 0, 0
        for row in rows:
            value = row[index]
            if missing(value):
                absent += 1
                continue
            try:
                if isinstance(value, bool):
                    raise InvalidOperation
                number = Decimal(str(value).strip())
                if not number.is_finite():
                    raise InvalidOperation
                values.append(number)
            except (InvalidOperation, ValueError):
                invalid += 1
        with localcontext() as context:
            # Avoid silently dropping small values next to very large decimals.
            context.prec = max(34, max((len(v.as_tuple().digits) + abs(v.as_tuple().exponent) + abs(v.adjusted()) for v in values), default=0) + len(str(len(values))) + 4)
            total = sum(values, Decimal(0)) if values else None
            average = total / len(values) if values else None
        report['numeric'][column] = {
            'valid': len(values), 'missing': absent, 'invalid': invalid,
            'sum': str(total) if total is not None else None,
            'mean': str(average) if average is not None else None,
            'min': str(min(values)) if values else None,
            'max': str(max(values)) if values else None,
        }
    return report


def main():
    parser = argparse.ArgumentParser(description='只读 CSV/Excel 数据质量检查')
    parser.add_argument('source', type=Path)
    parser.add_argument('--numeric', action='append', default=[])
    parser.add_argument('--encoding', default='utf-8-sig')
    parser.add_argument('--sheet')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    try:
        if not args.source.is_file() or args.source.is_symlink():
            raise ValueError('输入需为已核对的普通文件。')
        if args.output.exists() or args.output.is_symlink():
            raise ValueError('输出已存在，拒绝覆盖。')
        report = analyze(args.source, args.numeric, args.encoding, args.sheet)
        with args.output.open('x', encoding='utf-8') as handle:
            json.dump(report, handle, ensure_ascii=False, indent=2)
            handle.write('\n')
        print('质量报告已生成；尚未完成业务分析。')
    except (ValueError, OSError, UnicodeError, csv.Error) as error:
        parser.exit(2, f'数据检查未完成：{error}\n')


if __name__ == '__main__':
    main()
